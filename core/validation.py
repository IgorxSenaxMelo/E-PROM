"""Friendly validation of the input workbooks.

Instead of stopping at the first problem, these functions collect every issue
with its sheet and cell, so the user can fix the workbook in one pass.
"""

from __future__ import annotations
from pathlib import Path

import numpy as np
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from .data_loader import (
    _COMPARISON_KEYS,
    _CRITERIA_KEYS,
    _PRECISION_KEYS,
    _decision_maker_sheets,
    _find_evaluation_header,
    _find_fahp_header,
    _find_labelled_row,
    _is_empty,
    _text,
)
from .labels import (
    CONTINUOUS,
    ORDINAL,
    match_key,
    normalize_precision,
    normalize_type,
    parse_importance,
)


def _issue(severity, sheet, cell, message):
    return {"severity": severity, "sheet": sheet or "", "cell": cell or "", "message": message}


def _cell(r, c):
    return f"{get_column_letter(c)}{r}"


def validate_evaluation_workbook(path):
    """Return a list of issues found in the evaluation workbook."""
    issues = []
    try:
        wb = load_workbook(Path(path), data_only=True)
    except Exception as exc:
        return [_issue("error", "", "", f"The file could not be opened as .xlsx: {exc}")]

    sheets = _decision_maker_sheets(wb)
    if not sheets:
        return [_issue(
            "error", "", "",
            "No decision-maker sheet found (sheets named Guide, _Lists or "
            "Instructions are ignored).",
        )]

    reference = None
    for ws in sheets:
        header = _find_evaluation_header(ws)
        if header is None:
            issues.append(_issue(
                "error", ws.title, "A1",
                "Header row not found: expected 'ID' in column A and "
                "'Description' in column B.",
            ))
            continue

        columns, names = [], []
        for c in range(3, ws.max_column + 1):
            name = _text(ws.cell(header, c).value)
            if name:
                columns.append(c)
                names.append(name)
        if not names:
            issues.append(_issue("error", ws.title, _cell(header, 3), "No criterion columns found."))
            continue

        keys = [match_key(n) for n in names]
        for key in set(keys):
            if keys.count(key) > 1:
                issues.append(_issue("error", ws.title, "", f"Duplicated criterion name: '{names[keys.index(key)]}'."))

        ids = []
        n_rows = 0
        for r in range(header + 1, ws.max_row + 1):
            row = [ws.cell(r, c).value for c in range(1, ws.max_column + 1)]
            if all(_is_empty(v) for v in row):
                break
            n_rows += 1
            alt_id = _text(ws.cell(r, 1).value)
            if not alt_id:
                issues.append(_issue("error", ws.title, _cell(r, 1), "Missing alternative ID."))
            ids.append(alt_id)
            for c, name in zip(columns, names):
                value = ws.cell(r, c).value
                if _is_empty(value):
                    issues.append(_issue("error", ws.title, _cell(r, c), f"Empty value for '{name}' (alternative {alt_id or r})."))
                    continue
                try:
                    number = float(str(value).strip().replace(",", ".")) if not isinstance(value, (int, float)) else float(value)
                    if not np.isfinite(number):
                        raise ValueError
                except ValueError:
                    issues.append(_issue("error", ws.title, _cell(r, c), f"Non-numeric value {value!r} for '{name}'."))

        if n_rows < 2:
            issues.append(_issue("error", ws.title, "", "At least two alternatives are required."))
        non_empty = [i for i in ids if i]
        if len(set(non_empty)) != len(non_empty):
            issues.append(_issue("error", ws.title, "", "Duplicated alternative IDs."))

        current = {"keys": keys, "names": names, "ids": set(non_empty)}
        if reference is None:
            reference = (ws.title, current)
        else:
            ref_title, ref = reference
            if set(current["keys"]) != set(ref["keys"]):
                missing = [n for n, k in zip(ref["names"], ref["keys"]) if k not in current["keys"]]
                extra = [n for n, k in zip(names, keys) if k not in ref["keys"]]
                detail = []
                if missing:
                    detail.append("missing: " + ", ".join(missing))
                if extra:
                    detail.append("extra: " + ", ".join(extra))
                issues.append(_issue("error", ws.title, "", f"Criteria differ from sheet '{ref_title}' ({'; '.join(detail)})."))
            if current["ids"] != ref["ids"]:
                issues.append(_issue("error", ws.title, "", f"Alternative IDs differ from sheet '{ref_title}'."))

    if len(sheets) == 1:
        issues.append(_issue("warning", sheets[0].title, "", "Only one decision maker: group-agreement and leave-one-out analyses will not be available."))
    return issues


def validate_fahp_workbook(path, criteria=None, expert_names=None):
    """Return a list of issues found in the FAHP-Express workbook."""
    issues = []
    try:
        wb = load_workbook(Path(path), data_only=True)
    except Exception as exc:
        return [_issue("error", "", "", f"The file could not be opened as .xlsx: {exc}")]

    sheets = _decision_maker_sheets(wb)
    if not sheets:
        return [_issue("error", "", "", "No decision-maker sheet found in the FAHP-Express workbook.")]

    if expert_names is not None:
        if len(sheets) != len(expert_names):
            issues.append(_issue(
                "error", "", "",
                f"The FAHP-Express workbook has {len(sheets)} decision-maker sheets but the evaluation workbook has {len(expert_names)}.",
            ))
        elif {match_key(ws.title) for ws in sheets} != {match_key(n) for n in expert_names}:
            issues.append(_issue(
                "warning", "", "",
                "Sheet names differ between the two workbooks; decision makers will be paired by sheet position. Rename the sheets identically to pair them by name.",
            ))

    references = []
    for ws in sheets:
        header = _find_fahp_header(ws)
        if header is None:
            issues.append(_issue("error", ws.title, "A1", "Header row not found: expected 'Criteria' in column A followed by the criterion names."))
            continue

        columns, names = {}, []
        for c in range(2, ws.max_column + 1):
            name = _text(ws.cell(header, c).value)
            if name:
                columns[match_key(name)] = c
                names.append(name)

        if criteria is not None:
            wanted = {match_key(c): c for c in criteria}
            missing = [n for k, n in wanted.items() if k not in columns]
            extra = [n for n in names if match_key(n) not in wanted]
            if missing or extra:
                detail = []
                if missing:
                    detail.append("missing: " + ", ".join(missing))
                if extra:
                    detail.append("extra: " + ", ".join(extra))
                issues.append(_issue("error", ws.title, _cell(header, 1), f"Criteria differ from the evaluation workbook ({'; '.join(detail)})."))
                continue

        comp_row = _find_labelled_row(ws, header, _COMPARISON_KEYS, 1)
        prec_row = _find_labelled_row(ws, header, _PRECISION_KEYS, 2)
        if comp_row is None:
            issues.append(_issue("error", ws.title, "", "Row 'Comparison with reference' not found."))
            continue

        values = {}
        for name in names:
            col = columns[match_key(name)]
            raw = ws.cell(comp_row, col).value
            try:
                values[name] = parse_importance(raw)
            except ValueError as exc:
                issues.append(_issue("error", ws.title, _cell(comp_row, col), f"'{name}': {exc}"))

            if prec_row is not None:
                try:
                    normalize_precision(ws.cell(prec_row, col).value)
                except ValueError as exc:
                    issues.append(_issue("error", ws.title, _cell(prec_row, col), f"'{name}': {exc}"))
            elif name == names[0]:
                issues.append(_issue("warning", ws.title, "", "Row 'Attribute precision' not found: every criterion will use High precision."))

        if len(values) == len(names):
            ref_name = names[0]
            if not np.isclose(values[ref_name], 1.0):
                issues.append(_issue("error", ws.title, _cell(comp_row, columns[match_key(ref_name)]), f"The first criterion ('{ref_name}') is the reference and must be 'Equal importance' (1)."))
            low = [n for n, v in values.items() if v < 1.0 - 1e-9]
            if low:
                issues.append(_issue("error", ws.title, _cell(comp_row, columns[match_key(low[0])]), "Values below 1: the reference must be the most important criterion. Move the most important criterion to the first column. Affected: " + ", ".join(low)))
            vals = np.array(list(values.values()), dtype=float)
            if vals.min() > 0 and vals.max() / vals.min() > 9 + 1e-9:
                issues.append(_issue("error", ws.title, "", "The ratio between two criteria would exceed 9 (the Saaty limit); reduce the largest value."))
            references.append(ref_name)

    if len({match_key(r) for r in references}) > 1:
        issues.append(_issue("info", "", "", "Decision makers use different reference criteria (" + ", ".join(sorted(set(references))) + "). This is supported."))
    return issues


def validate_against_config(data, types, scales):
    """Check the evaluation values against the configured data types / scales."""
    issues = []
    for e, name in enumerate(data["expert_names"]):
        matrix = np.asarray(data["evaluations"][name], dtype=float)
        for k, criterion in enumerate(data["criteria"]):
            typ = normalize_type(types[k])
            col = matrix[:, k]
            if typ == ORDINAL:
                top = int(scales[k])
                for i, v in enumerate(col):
                    if abs(v - round(v)) > 1e-9 or v < 1 or v > top:
                        issues.append(_issue(
                            "error", name, "",
                            f"'{criterion}', alternative {data['ids'][i]}: value {v:g} is not an integer in the 1-{top} ordinal scale.",
                        ))
            else:
                if np.any(col == 0):
                    issues.append(_issue("warning", name, "", f"'{criterion}': zero values have no fuzzy width (precision has no effect on them)."))
                if np.any(col < 0):
                    issues.append(_issue("warning", name, "", f"'{criterion}': negative values found; check the preference direction."))
    return issues


def summarize(issues):
    return {
        "errors": sum(1 for i in issues if i["severity"] == "error"),
        "warnings": sum(1 for i in issues if i["severity"] == "warning"),
        "info": sum(1 for i in issues if i["severity"] == "info"),
    }
