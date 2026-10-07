from pathlib import Path
import io
import sys
import tempfile

import numpy as np
import pandas as pd
import streamlit as st
import plotly.graph_objects as go

ROOT = Path(__file__).resolve().parent
sys.path.insert(0, str(ROOT))

from core.data_loader_v73 import (
    read_expert_workbook,
    read_fahp_workbook,
)
from core.e_prom_core import run_e_prom
from core.fuzzy_numbers import defuzz_coa
from core.cpp import cpp_normal
from core.monte_carlo import run_e_prom_monte_carlo


st.set_page_config(
    page_title="E-PROM — Fuzzy Preference Ranking",
    page_icon="⚖️",
    layout="wide",
)


PREF_CODES = {
    "Usual": 1,
    "U-shape": 2,
    "V-shape": 3,
    "Level": 4,
    "Linear": 5,
}

# Spreadsheet input vocabulary. These are translated to the numeric values
# expected by the existing FAHP-Express core; the core itself is unchanged.
LINGUISTIC_SCALE = {
    "equal importance": 1,
    "weakly more important": 2,
    "moderately more important": 3,
    "moderately to strongly more important": 4,
    "strongly more important": 5,
    "strongly to very strongly more important": 6,
    "very strongly more important": 7,
    "very strongly to extremely more important": 8,
    "extremely more important": 9,
}

PRECISION_SCALE = {
    "high": "Alta",
    "medium": "Média",
    "low": "Baixa",
    # Backward compatibility with Portuguese workbooks.
    "alta": "Alta",
    "média": "Média",
    "media": "Média",
    "baixa": "Baixa",
}


def save_uploaded(f):
    tmp = tempfile.NamedTemporaryFile(
        delete=False,
        suffix=Path(f.name).suffix or ".xlsx",
    )
    tmp.write(f.getvalue())
    tmp.close()
    return Path(tmp.name)


def _criterion_key(value):
    """Normaliza apenas espaços externos para comparação de nomes."""
    if value is None:
        return ""
    return str(value).strip()


def _find_header_row(ws, required_headers, max_scan_rows=20):
    """Localiza a linha de cabeçalho pelo conjunto de nomes esperado."""
    required = {_criterion_key(x).casefold() for x in required_headers}
    for r in range(1, min(ws.max_row, max_scan_rows) + 1):
        values = {
            _criterion_key(ws.cell(r, c).value).casefold()
            for c in range(1, ws.max_column + 1)
        }
        if required.issubset(values):
            return r
    return None


def _find_fahp_header_row(ws, max_scan_rows=30):
    """Locate the FAHP criterion header without relying on sheet-title text.

    The standard English template uses:
        Criteria | criterion 1 | criterion 2 | ...

    Legacy Portuguese workbooks use:
        Critérios | critério 1 | critério 2 | ...

    This helper deliberately checks the actual header row and does not infer
    the header from A1, sheet name, or the order of criteria.
    """
    for r in range(1, min(ws.max_row, max_scan_rows) + 1):
        first = _criterion_key(ws.cell(r, 1).value).casefold()
        if first not in {"criteria", "critérios"}:
            continue

        # Require at least two criterion names to avoid accepting a title row.
        nonempty = 0
        for c in range(2, ws.max_column + 1):
            if _criterion_key(ws.cell(r, c).value):
                nonempty += 1

        if nonempty >= 2:
            return r

    return None


def _criterion_columns(ws, header_row, first_criterion_col):
    """Retorna {nome_do_criterio: coluna} a partir do cabeçalho."""
    result = {}
    for c in range(first_criterion_col, ws.max_column + 1):
        value = _criterion_key(ws.cell(header_row, c).value)
        if not value:
            continue
        if value in result:
            raise ValueError(
                f"Critério duplicado na aba '{ws.title}': '{value}'."
            )
        result[value] = c
    return result


def _reorder_columns_in_place(ws, header_row, first_criterion_col, canonical):
    """Reordena somente o bloco de critérios, preservando os valores associados."""
    col_map = _criterion_columns(ws, header_row, first_criterion_col)
    current = list(col_map.keys())

    canonical_keys = [_criterion_key(x) for x in canonical]
    current_keys = [_criterion_key(x) for x in current]

    if set(current_keys) != set(canonical_keys):
        missing = [x for x in canonical_keys if x not in current_keys]
        extra = [x for x in current_keys if x not in canonical_keys]
        details = []
        if missing:
            details.append("ausentes: " + ", ".join(missing))
        if extra:
            details.append("extras: " + ", ".join(extra))
        raise ValueError(
            f"Aba '{ws.title}' possui critérios diferentes dos demais especialistas "
            f"({'; '.join(details)})."
        )

    # Nenhuma alteração é feita quando a ordem já é a mesma.
    if current_keys == canonical_keys:
        return

    # O aplicativo trabalha com os valores das células. Fazemos a reordenação
    # em uma cópia dos valores para que cada avaliação/referência permaneça
    # ligada ao nome do critério, independentemente da posição da coluna.
    max_row = ws.max_row
    snapshots = {
        name: [ws.cell(r, col_map[name]).value for r in range(header_row, max_row + 1)]
        for name in canonical_keys
    }

    target_cols = [col_map[name] for name in current_keys]
    target_cols.sort()

    for target_col, name in zip(target_cols, canonical_keys):
        values = snapshots[name]
        for offset, value in enumerate(values):
            ws.cell(header_row + offset, target_col).value = value


def translate_fahp_inputs(ws, header_row, col_map):
    # Find the two semantic rows by their labels, independently of
    # criterion order.
    comparison_row = None
    precision_row = None

    for r in range(header_row, ws.max_row + 1):
        label = ws.cell(r, 1).value
        if label is None:
            continue
        key = str(label).strip().casefold()
        if key == "comparação com referência" or key == "comparison with reference":
            comparison_row = r
        elif key == "precisão do atributo" or key == "attribute precision":
            precision_row = r

    if comparison_row is not None:
        for col in col_map.values():
            value = ws.cell(comparison_row, col).value
            if isinstance(value, str):
                key = " ".join(value.strip().casefold().split())
                if key in LINGUISTIC_SCALE:
                    ws.cell(comparison_row, col).value = LINGUISTIC_SCALE[key]
                elif key:
                    # Preserve numeric strings for backward compatibility.
                    try:
                        number = float(value.replace(",", "."))
                        if number.is_integer():
                            number = int(number)
                        ws.cell(comparison_row, col).value = number
                    except ValueError:
                        pass

    if precision_row is not None:
        for col in col_map.values():
            value = ws.cell(precision_row, col).value
            if isinstance(value, str):
                key = " ".join(value.strip().casefold().split())
                if key in PRECISION_SCALE:
                    ws.cell(precision_row, col).value = PRECISION_SCALE[key]

def normalize_expert_workbook(path):
    """
    Aceita especialistas com os mesmos critérios em ordens diferentes.
    A primeira aba de especialista define a ordem canônica; nas demais,
    as colunas são realinhadas pelo nome do critério, sem alterar os valores.
    """
    from openpyxl import load_workbook

    # Linguistic FAHP-Express scale.
    # The core remains numeric; only the spreadsheet input layer translates
    # the linguistic labels into the existing 1-9 Saaty-compatible values.
    linguistic_scale = {
        "equal importance": 1,
        "weakly more important": 2,
        "moderately more important": 3,
        "moderately to strongly more important": 4,
        "strongly more important": 5,
        "strongly to very strongly more important": 6,
        "very strongly more important": 7,
        "very strongly to extremely more important": 8,
        "extremely more important": 9,
    }

    precision_scale = {
        "high": "Alta",
        "medium": "Média",
        "low": "Baixa",
        # Backward compatibility with existing Portuguese spreadsheets.
        "alta": "Alta",
        "média": "Média",
        "media": "Média",
        "baixa": "Baixa",
    }



    wb = load_workbook(path, data_only=False)
    sheets = [
        ws for ws in wb.worksheets
        if _criterion_key(ws.title).casefold() not in {"guide", "_lists"}
    ]

    if not sheets:
        raise ValueError("A planilha de avaliações não possui abas de especialistas.")

    first = sheets[0]
    header_row = _find_header_row(first, ["ID", "Descrição"])
    if header_row is None:
        raise ValueError(
            f"Não foi possível localizar o cabeçalho de avaliações na aba '{first.title}'."
        )

    first_cols = _criterion_columns(first, header_row, 3)
    canonical = list(first_cols.keys())
    if not canonical:
        raise ValueError(
            f"Nenhum critério foi encontrado na aba '{first.title}'."
        )

    for ws in sheets:
        row = _find_header_row(ws, ["ID", "Descrição"])
        if row is None:
            raise ValueError(
                f"Não foi possível localizar o cabeçalho de avaliações na aba '{ws.title}'."
            )
        _reorder_columns_in_place(ws, row, 3, canonical)

    out = tempfile.NamedTemporaryFile(delete=False, suffix=".xlsx")
    out.close()
    wb.save(out.name)
    return Path(out.name)


def normalize_fahp_workbook(path, expected_criteria):
    """
    Normaliza a planilha FAHP-Express pela identidade dos critérios, mas
    preserva uma referência específica para cada decisor.

    A referência de cada aba é o primeiro critério originalmente apresentado
    naquela aba, conforme a convenção utilizada pelo data_loader_v73.py.
    Portanto, decisores diferentes podem escolher referências diferentes.

    A ordem final entregue ao data_loader é:
        [referência do decisor] + [demais critérios na ordem canônica]

    Os valores de "Comparação com referência" e "Precisão do atributo"
    permanecem associados ao respectivo critério.
    """
    from openpyxl import load_workbook
    import unicodedata
    import re

    def match_key(value):
        """Chave de comparação tolerante a maiúsculas, acentos e separadores."""
        text = _criterion_key(value).casefold()
        text = unicodedata.normalize("NFKD", text)
        text = "".join(
            ch for ch in text
            if not unicodedata.combining(ch)
        )
        text = re.sub(r"[-–—_/]+", " ", text)
        text = re.sub(r"\\s+", " ", text).strip()

        # Compatibilidade com a nomenclatura usada na planilha FAHP antiga:
        # "CUSTO" corresponde a "Custo-benefício".
        if text in {"custo", "custo beneficio"}:
            return "custo-beneficio"

        return text

    wb = load_workbook(path, data_only=False)
    sheets = [
        ws for ws in wb.worksheets
        if _criterion_key(ws.title).casefold() not in {"guide", "_lists"}
    ]

    if not sheets:
        raise ValueError(
            "A planilha FAHP-Express não possui abas de especialistas."
        )

    canonical = [_criterion_key(x) for x in expected_criteria]
    if not canonical or any(not x for x in canonical):
        raise ValueError(
            "Não foi possível determinar os critérios da planilha de avaliações."
        )

    canonical_by_key = {}
    for name in canonical:
        key = match_key(name)
        if key in canonical_by_key:
            raise ValueError(
                f"Critérios ambíguos na planilha de avaliações: "
                f"'{canonical_by_key[key]}' e '{name}'."
            )
        canonical_by_key[key] = name

    canonical_keys = [match_key(name) for name in canonical]

    for ws in sheets:
        # The standard template places the criterion header on the row
        # whose first cell is "Criteria" (English) or "Critérios" (legacy).
        # Do not depend on the title text in A1; it is only a sheet title.
        header_row = _find_fahp_header_row(ws)
        if header_row is None:
            raise ValueError(
                f"Could not locate the FAHP-Express criteria header "
                f"in sheet '{ws.title}'. Expected a row beginning with "
                f"'Criteria' or 'Critérios'."
            )

        # Normalize English workbook labels back to the legacy labels expected
        # by data_loader_v73.py. This keeps the computational core unchanged.
        internal_labels = {
            "Criteria": "Critérios",
            "Comparison with reference": "Comparação com referência",
            "Attribute precision": "Precisão do atributo",
            "Reference": "Referência",
            "Interpretation": "Interpretação",
            "Notes": "Observação",
        }
        for r in range(1, ws.max_row + 1):
            value = ws.cell(r, 1).value
            if isinstance(value, str) and value.strip() in internal_labels:
                ws.cell(r, 1).value = internal_labels[value.strip()]

        col_map = _criterion_columns(ws, header_row, 2)
        current_names = list(col_map.keys())
        current_keys = [match_key(x) for x in current_names]

        if len(current_keys) != len(set(current_keys)):
            raise ValueError(
                f"Aba FAHP '{ws.title}' possui critérios duplicados."
            )

        current_by_key = {
            key: name
            for key, name in zip(current_keys, current_names)
        }

        missing = [
            canonical_by_key[key]
            for key in canonical_keys
            if key not in current_by_key
        ]
        extra = [
            current_by_key[key]
            for key in current_keys
            if key not in canonical_by_key
        ]

        if missing or extra:
            details = []
            if missing:
                details.append("ausentes: " + ", ".join(missing))
            if extra:
                details.append("extras: " + ", ".join(extra))
            raise ValueError(
                f"Aba FAHP '{ws.title}' possui critérios diferentes "
                f"({'; '.join(details)})."
            )

        # Translate the spreadsheet representation only at the input layer.
        # The core continues to receive the same numeric representation.
        translate_fahp_inputs(ws, header_row, col_map)

        # A referência é individual de cada decisor: é o primeiro critério
        # que estava na aba original. Não usamos a referência da primeira aba.
        reference_key = current_keys[0]

        # O data_loader exige a referência na primeira coluna.
        target_keys = [reference_key] + [
            key for key in canonical_keys
            if key != reference_key
        ]

        if current_keys == target_keys:
            continue

        # Captura os valores por identidade do critério antes de escrever.
        snapshots = {
            key: [
                ws.cell(r, col_map[current_by_key[key]]).value
                for r in range(header_row, ws.max_row + 1)
            ]
            for key in current_keys
        }

        target_cols = sorted(col_map.values())

        for target_col, key in zip(target_cols, target_keys):
            values = snapshots[key]
            for offset, value in enumerate(values):
                ws.cell(
                    header_row + offset,
                    target_col,
                ).value = value

    out = tempfile.NamedTemporaryFile(delete=False, suffix=".xlsx")
    out.close()
    wb.save(out.name)
    return Path(out.name)


def result_df(data, result):
    return pd.DataFrame({
        "Rank": np.asarray(
            result["rank_e_prom"],
            dtype=int,
        ),
        "Alternative": data["ids"],
        "Description": data["descriptions"],
        "Phi+": result["phi_plus"],
        "Phi-": result["phi_minus"],
        "Net Phi": result["phi_net"],
    }).sort_values(
        ["Rank", "Alternative"]
    ).reset_index(drop=True)


def excel_bytes(data, result, config, fahp):
    out = io.BytesIO()

    with pd.ExcelWriter(
        out,
        engine="openpyxl",
    ) as writer:

        result_df(
            data,
            result,
        ).to_excel(
            writer,
            sheet_name="Ranking",
            index=False,
        )

        phi = pd.DataFrame(
            result["phi_net_experts"],
            index=data["ids"],
            columns=result["expert_names"],
        )
        phi.index.name = "Alternative"

        phi.to_excel(
            writer,
            sheet_name="Phi_by_Decision_Maker",
        )

        config.to_excel(
            writer,
            sheet_name="Configuration",
            index=False,
        )

        fuzzy_w = pd.DataFrame(
            fahp["w_fuzzy_group"],
            index=data["criteria"],
            columns=["a", "b", "c", "d"],
        )
        fuzzy_w.index.name = "Criterion"

        fuzzy_w.to_excel(
            writer,
            sheet_name="Fuzzy_Weights",
        )

        precision_matrix = fahp.get(
            "precision_levels",
            [["Alta"] * len(data["criteria"]) for _ in fahp.get("expert_names", [])],
        )
        precision_df = pd.DataFrame(
            precision_matrix,
            index=fahp.get("expert_names", []),
            columns=data["criteria"],
        )
        precision_df = precision_df.replace({
            "Alta": "High",
            "Média": "Medium",
            "Baixa": "Low",
        })
        precision_df.index.name = "Decision maker"
        precision_df.to_excel(
            writer,
            sheet_name="Attribute_Precision",
        )

    out.seek(0)
    return out.getvalue()


def plot_weights(fahp, criteria):
    wf = np.asarray(
        fahp["w_fuzzy_individual"],
        dtype=float,
    )

    crisp = np.array([
        [
            defuzz_coa(wf[k, :, e])
            for e in range(wf.shape[2])
        ]
        for k in range(wf.shape[0])
    ])

    mean = crisp.mean(axis=1)

    fig = go.Figure()

    fig.add_trace(
        go.Bar(
            x=criteria,
            y=mean,
            name="Mean weight",
        )
    )

    for e, name in enumerate(
        fahp.get(
            "expert_names",
            fahp.get(
                "sheets",
                [f"E{i + 1}" for i in range(crisp.shape[1])],
            ),
        )
    ):
        fig.add_trace(
            go.Scatter(
                x=criteria,
                y=crisp[:, e],
                mode="markers",
                name=name,
            )
        )

    fig.update_layout(
        title="Fuzzy criterion weights — FAHP-Express",
        xaxis_title="Criterion",
        yaxis_title="Weight",
        height=500,
    )

    return fig


def relation_matrix(result):
    plus = np.asarray(
        result["phi_plus"],
        dtype=float,
    )
    minus = np.asarray(
        result["phi_minus"],
        dtype=float,
    )

    n = len(plus)
    rel = np.empty(
        (n, n),
        dtype=object,
    )

    eps = 1e-10

    for i in range(n):
        for j in range(n):

            if i == j:
                rel[i, j] = "I"
                continue

            a = (
                plus[i] > plus[j] + eps
                and minus[i] < minus[j] - eps
            )

            b = (
                plus[i] < plus[j] - eps
                and minus[i] > minus[j] + eps
            )

            eq = (
                abs(plus[i] - plus[j]) <= eps
                and abs(minus[i] - minus[j]) <= eps
            )

            if a:
                rel[i, j] = "P"
            elif b:
                rel[i, j] = "P-"
            elif eq:
                rel[i, j] = "I"
            else:
                rel[i, j] = "R"

    return rel


def alternative_labels(data):
    """Retorna os nomes das alternativas para exibição nos gráficos.

    Os cálculos continuam usando os IDs internamente. Quando houver uma
    descrição válida, ela é usada apenas como rótulo visual.
    """
    ids = [str(x) for x in data["ids"]]
    descriptions = data.get("descriptions", [])

    if descriptions is None or len(descriptions) != len(ids):
        return ids

    labels = []
    used = {}

    for alt_id, desc in zip(ids, descriptions):
        label = str(desc).strip() if desc is not None else ""
        if not label:
            label = alt_id

        # Evita rótulos visualmente indistinguíveis quando duas
        # alternativas possuem a mesma descrição.
        count = used.get(label, 0) + 1
        used[label] = count
        if count > 1:
            label = f"{label} ({alt_id})"

        labels.append(label)

    return labels


def plot_relation(data, result):
    rel = relation_matrix(result)

    mapping = {
        "P-": 0,
        "I": 1,
        "P": 2,
        "R": 3,
    }

    z = np.vectorize(
        mapping.get
    )(rel)

    labels = alternative_labels(data)

    fig = go.Figure(
        go.Heatmap(
            z=z,
            x=labels,
            y=labels,
            text=rel,
            texttemplate="%{text}",
            zmin=0,
            zmax=3,
            colorscale=[
                [0.00, "#3333cc"],
                [0.249, "#3333cc"],
                [0.250, "#aaaaaa"],
                [0.499, "#aaaaaa"],
                [0.500, "#22aa22"],
                [0.749, "#22aa22"],
                [0.750, "#e8cc28"],
                [1.00, "#e8cc28"],
            ],
            colorbar=dict(
                tickvals=[0, 1, 2, 3],
                ticktext=["P-", "I", "P", "R"],
                title="Relation",
            ),
        )
    )

    fig.update_layout(
        title="PROMETHEE I — relations",
        height=max(750, len(labels) * 110),
        margin=dict(l=180, r=80, t=80, b=120),
    )

    return fig


def plot_rank_heatmap(data, result):
    labels = alternative_labels(data)

    df = pd.DataFrame(
        result["rank_experts"],
        index=labels,
        columns=result["expert_names"],
    )

    df["E-PROM"] = result["rank_e_prom"]

    df = df.sort_values(
        "E-PROM"
    )

    fig = go.Figure(
        go.Heatmap(
            z=df.values,
            x=list(df.columns),
            y=list(df.index),
            text=np.vectorize(
                lambda x: f"{x:g}"
            )(df.values),
            texttemplate="%{text}",
            colorscale="Blues_r",
            zmin=1,
            zmax=len(labels),
            xgap=1,
            ygap=1,
            colorbar=dict(
                title="Position"
            ),
        )
    )

    fig.update_layout(
        title="Positions by decision maker and E-PROM",
        height=max(600, len(labels) * 55),
        margin=dict(l=180, r=80, t=80, b=100),
    )

    return fig



def plot_mc_rank_acceptability(data, mc):
    labels = alternative_labels(data)
    order = np.argsort(mc["mean_rank"])
    M = mc["rank_acceptability"][order, :]

    fig = go.Figure(
        go.Heatmap(
            z=M,
            x=list(range(1, len(labels) + 1)),
            y=[labels[i] for i in order],
            text=np.vectorize(lambda x: f"{x:.1f}%")(M),
            texttemplate="%{text}",
            colorscale="Blues",
            zmin=0,
            zmax=max(1, np.nanmax(M)),
            colorbar=dict(title="%"),
        )
    )
    fig.update_layout(
        title="CPP/Monte Carlo — rank acceptability",
        xaxis_title="Position",
        yaxis_title="Alternative",
        height=max(550, len(labels) * 55),
        margin=dict(l=180, r=80, t=80, b=100),
    )
    return fig


def plot_mc_outworking(data, mc):
    labels = alternative_labels(data)
    order = np.argsort(-mc["outranking_acceptability"].sum(axis=1))
    M = mc["outranking_acceptability"][np.ix_(order, order)]

    fig = go.Figure(
        go.Heatmap(
            z=M,
            x=[labels[i] for i in order],
            y=[labels[i] for i in order],
            colorscale="Blues",
            zmin=0,
            zmax=100,
            colorbar=dict(title="%"),
        )
    )
    fig.update_layout(
        title="CPP/Monte Carlo — outranking acceptability",
        xaxis_title="Outranked alternative",
        yaxis_title="Outranking alternative",
        height=max(700, len(labels) * 55),
        margin=dict(l=180, r=120, t=100, b=180),
    )
    return fig


def mc_summary_df(data, mc):
    return pd.DataFrame({
        "Alternative": data["ids"],
        "Mean position": mc["mean_rank"],
        "Position SD": mc["std_rank"],
        "1st-place frequency": 100 * mc["freq_top1"],
        "Mean Phi": mc["mean_phi"],
        "Phi SD": mc["std_phi"],
    }).sort_values(
        "Mean position"
    ).reset_index(drop=True)

# ================================================================
# HEADER
# ================================================================

st.title("E-PROM")
st.subheader(
    "Express Fuzzy Preference Ranking with PROMETHEE — v7.3"
)
st.caption(
    "FAHP-Express + Fuzzy PROMETHEE — generalized method."
)

# ================================================================
# MODELOS PARA DOWNLOAD
# ================================================================

def make_english_linguistic_fahp_template(template_bytes):
    """Convert the existing FAHP-Express template to the English linguistic-input form.

    This is presentation/input-layer conversion only. The workbook used for
    computation is converted back to the numeric representation by
    normalize_fahp_workbook().
    """
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(template_bytes), data_only=False)

    linguistic_terms = {
        1: "Equal importance",
        2: "Weakly more important",
        3: "Moderately more important",
        4: "Moderately to strongly more important",
        5: "Strongly more important",
        6: "Strongly to very strongly more important",
        7: "Very strongly more important",
        8: "Very strongly to extremely more important",
        9: "Extremely more important",
    }

    for ws in wb.worksheets:
        if str(ws.title).strip().casefold() == "guide":
            # Translate the basic guide labels where they exist.
            for row in ws.iter_rows():
                for cell in row:
                    if isinstance(cell.value, str):
                        cell.value = {
                            "Campo": "Field",
                            "Orientação": "Guidance",
                            "Referência": "Reference",
                            "Demais critérios": "Other criteria",
                            "Precisão": "Precision",
                            "Critérios": "Criteria",
                            "Decisores": "Decision makers",
                            "Romance é o critério de referência e recebe comparação 1 consigo mesmo.":
                                "The reference criterion receives Equal importance (1) against itself.",
                            "Fator de 1 a 9 indica quanto a referência é mais importante que o critério correspondente.":
                                "The 1–9 linguistic scale indicates how much more important the reference is than the corresponding criterion.",
                            "Alta, Média ou Baixa; controla a largura da representação fuzzy.":
                                "High, Medium, or Low; controls fuzzy representation width.",
                            "Somente Decisor 1 e Decisor 2 são utilizados.":
                                "Only Decision maker 1 and Decision maker 2 are used.",
                        }.get(cell.value, cell.value)
            continue

        # Translate labels in the first column.
        for r in range(1, ws.max_row + 1):
            label = ws.cell(r, 1).value
            if isinstance(label, str):
                label_map = {
                    "Critérios": "Criteria",
                    "Comparação com referência": "Comparison with reference",
                    "Precisão do atributo": "Attribute precision",
                    "Referência": "Reference",
                    "Interpretação": "Interpretation",
                    "Observação": "Notes",
                }
                ws.cell(r, 1).value = label_map.get(label, label)

        # Replace numeric 1–9 comparison values with linguistic terms.
        for r in range(1, ws.max_row + 1):
            label = ws.cell(r, 1).value
            if str(label).strip().casefold() == "comparison with reference":
                for c in range(2, ws.max_column + 1):
                    value = ws.cell(r, c).value
                    if isinstance(value, (int, float)) and int(value) in linguistic_terms:
                        ws.cell(r, c).value = linguistic_terms[int(value)]

            if str(label).strip().casefold() == "attribute precision":
                for c in range(2, ws.max_column + 1):
                    value = ws.cell(r, c).value
                    if isinstance(value, str):
                        pmap = {"Alta": "High", "Média": "Medium", "Baixa": "Low",
                                "High": "High", "Medium": "Medium", "Low": "Low"}
                        ws.cell(r, c).value = pmap.get(value.strip(), value)

    out = io.BytesIO()
    wb.save(out)
    out.seek(0)
    return out.getvalue()


TEMPLATES_DIR = ROOT / "templates"
EVAL_TEMPLATE = TEMPLATES_DIR / "E_PROM_Avaliacoes_Template.xlsx"
FAHP_TEMPLATE = TEMPLATES_DIR / "E_PROM_FAHP_Express_Template.xlsx"

st.markdown("### 📥 Download the template workbooks")
st.caption(
    "Use these files as a starting point. Enter the experts’ evaluations and, separately, the FAHP-Express references."
)

col_model_1, col_model_2 = st.columns(2)

with col_model_1:
    if EVAL_TEMPLATE.exists():
        with open(EVAL_TEMPLATE, "rb") as f:
            st.download_button(
                "⬇️ Evaluation Workbook",
                data=f.read(),
                file_name=EVAL_TEMPLATE.name,
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                use_container_width=True,
                key="download_eval_template",
            )
    else:
        st.warning("Evaluation template not found in the package.")

with col_model_2:
    if FAHP_TEMPLATE.exists():
        with open(FAHP_TEMPLATE, "rb") as f:
            template_bytes = make_english_linguistic_fahp_template(f.read())
            st.download_button(
                "⬇️ FAHP-Express Workbook",
                data=template_bytes,
                file_name="E_PROM_FAHP_Express_Linguistic_Template.xlsx",
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                use_container_width=True,
                key="download_fahp_template",
            )
    else:
        st.warning("FAHP-Express template not found in the package.")

st.divider()

GUIDE_TEXT = """# E-PROM — Guide

## 1. Objective

The E-PROM (Express Fuzzy Preference Ranking with PROMETHEE — v7.3) combines:

**FAHP-Express → fuzzy criterion weights → Fuzzy PROMETHEE**

The method keeps fuzzy weights during Fuzzy PROMETHEE and applies defuzzification
to the final flows.

---

## 2. Evaluation file

The evaluation workbook must contain:

- one sheet per decision maker;
- a final sheet named `Guide`, which is not processed;
- column A: alternative ID;
- column B: description;
- column C onward: criteria.

All decision makers must use the same alternative IDs, descriptions, and criterion names.
The order of criteria may differ between decision makers; the application aligns them
by criterion name while preserving each decision maker's own reference criterion.

---

## 3. Criterion type

### Ordinal

Use when the evaluation is an ordered scale, for example:

- 1–5;
- 1–7;
- 1–9.

Allowed preference functions:

- **Usual**
- **Level**

The ordinal value is converted to a trapezoidal fuzzy number.

### Continuous

Use for quantities such as:

- cost;
- schedule;
- distance;
- time;
- consumption;
- physical performance.

Allowed preference functions:

- **U-shape**
- **V-shape**
- **Linear**

The value remains in its original unit.

---

## 4. FAHP-Express linguistic input

The FAHP-Express workbook accepts linguistic comparison terms instead of numeric
comparison factors:

| Value | Linguistic term |
|---:|---|
| 1 | Equal importance |
| 2 | Weakly more important |
| 3 | Moderately more important |
| 4 | Moderately to strongly more important |
| 5 | Strongly more important |
| 6 | Strongly to very strongly more important |
| 7 | Very strongly more important |
| 8 | Very strongly to extremely more important |
| 9 | Extremely more important |

Each decision maker may use a different reference criterion. The reference
criterion receives **Equal importance (1)**.

The application translates these terms internally to the numeric 1–9 scale
before calling the existing FAHP-Express core. The mathematical core is unchanged.

### Attribute precision

Use:

- **High**
- **Medium**
- **Low**

These labels are converted internally to the existing precision representation.

---

## 5. Reference criterion

Each decision maker may select a different reference criterion.

For example:

- Decision maker 1 → Romance;
- Decision maker 2 → Memorable experiences.

The reference criterion does not need to be the same across decision makers.

The application identifies the reference from the first criterion in each FAHP-Express
sheet, as required by the existing data loader, while preserving the criterion identity.

---

## 6. Fuzzy PROMETHEE

The E-PROM does not defuzzify the criterion weights before PROMETHEE.

For each criterion:

`fuzzy weight × fuzzy preference`

The terms are aggregated to obtain fuzzy flows:

- Φ+;
- Φ−;
- net Φ.

Defuzzification by center of area is applied to the final flows to obtain the ranking.

---

## 7. Monte Carlo / CPP

The Monte Carlo module evaluates the stability of the result while preserving
the fuzzy information during expert aggregation.

### Ordinal criteria

The distribution is empirical, constructed directly from the experts' responses.

### Continuous criteria

The user specifies a percentage variation.

### Weights

At each iteration, the individual FAHP-Express reference rows are resampled
and a new set of fuzzy weights is calculated.

---

## 8. Main Monte Carlo results

The E-PROM calculates:

- mean position;
- standard deviation of position;
- mean net Φ;
- standard deviation of net Φ;
- first-place frequency;
- rank acceptability matrix;
- PROMETHEE I outranking acceptability matrix;
- criterion-weight distributions.

---

## 9. Interpretation

The probabilistic analysis should be used as a stability analysis, not as an
automatic replacement for the deterministic ranking.

Useful results include:

- frequency with which an alternative occupies each position;
- first-place frequency;
- Φ dispersion;
- weight sensitivity;
- pairwise outranking probability.

---

## 10. Recommended workflow

1. Prepare the decision makers' evaluations.
2. Prepare the FAHP-Express workbook using linguistic comparison terms.
3. Configure criterion type and direction.
4. Select the preference function.
5. Define q and/or p.
6. Run deterministic E-PROM.
7. Check the ranking and PROMETHEE I relations.
8. Run Monte Carlo/CPP.
9. Evaluate stability and position acceptability.
10. Export the results.
"""



# ================================================================
# SIDEBAR
# ================================================================

with st.sidebar:

    st.header("1. Inputs")

    req_file = st.file_uploader(
        "Evaluation workbook",
        type=["xlsx"],
    )

    fahp_file = st.file_uploader(
        "FAHP-Express workbook",
        type=["xlsx"],
    )

    run = False

    if req_file is not None:

        try:

            preview = read_expert_workbook(
                normalize_expert_workbook(save_uploaded(req_file))
            )

            criteria = preview["criteria"]
            ncrit = len(criteria)

            st.success(
                f"{preview['n_experts']} decision makers · "
                f"{preview['n_requirements']} alternatives · "
                f"{ncrit} criteria"
            )

            st.divider()
            st.header("2. Criterion configuration")

            st.info(
                "Ordinal: Usual or Level. Continuous: U-shape, V-shape, or Linear."
            )

            dirs = []
            types = []
            prefs = []
            qs = []
            ps = []
            scales = []
            rows = []

            for k, criterion in enumerate(
                criteria
            ):

                st.markdown(
                    f"### {criterion}"
                )

                typ_display = st.selectbox(
                    "Data type",
                    ["Ordinal", "Continuous"],
                    key=f"type_{k}",
                )
                # Preserve the existing internal representation used by the core.
                typ = "Ordinal" if typ_display == "Ordinal" else "Contínuo"

                direction_name = st.selectbox(
                    "Direction",
                    ["Maximize", "Minimize"],
                    key=f"direction_{k}",
                )

                if typ == "Ordinal":

                    scale = st.selectbox(
                        "Ordinal scale",
                        [5, 7, 9],
                        index=1,
                        key=f"scale_{k}",
                    )

                    pref_name = st.selectbox(
                        "Preference function",
                        ["Usual", "Level"],
                        key=f"pref_{k}",
                    )

                    if pref_name == "Usual":
                        q = 0.0
                        p = 0.0
                    else:

                        q = st.number_input(
                            "q — indifference",
                            min_value=0.0,
                            value=0.0,
                            step=1.0,
                            key=f"q_{k}",
                        )

                        p = st.number_input(
                            "p — preference",
                            min_value=float(q + 1e-9),
                            value=float(q + 1),
                            step=1.0,
                            key=f"p_{k}",
                        )

                else:

                    scale = 0
                    st.caption(
                        "Continuous values are represented by a trapezoidal fuzzy number. "
                        "na unidade original. A largura do trapezoide é definida "
                        "pela precisão informada na planilha FAHP-Express."
                    )

                    pref_name = st.selectbox(
                        "Preference function",
                        [
                            "U-shape",
                            "V-shape",
                            "Linear",
                        ],
                        key=f"pref_{k}",
                    )

                    if pref_name == "U-shape":

                        q = st.number_input(
                            "q — indifference threshold",
                            min_value=0.0,
                            value=0.0,
                            step=1.0,
                            key=f"q_{k}",
                            help=(
                                "Diferença até q: indiferença. "
                                "Use a unidade original do critério."
                            ),
                        )

                        p = 0.0

                    elif pref_name == "V-shape":

                        q = 0.0

                        p = st.number_input(
                            "p — preference threshold",
                            min_value=1e-9,
                            value=1.0,
                            step=1.0,
                            key=f"p_{k}",
                            help=(
                                "Até p, a preferência cresce linearmente; "
                                "a partir de p, é completa."
                            ),
                        )

                    else:  # Linear

                        q = st.number_input(
                            "q — indifference threshold",
                            min_value=0.0,
                            value=0.0,
                            step=1.0,
                            key=f"q_{k}",
                        )

                        p = st.number_input(
                            "p — preference threshold",
                            min_value=float(q + 1e-9),
                            value=float(q + 1),
                            step=1.0,
                            key=f"p_{k}",
                        )

                code = PREF_CODES[pref_name]

                dirs.append(
                    1
                    if direction_name == "Maximize"
                    else -1
                )

                types.append(typ)
                prefs.append(code)
                qs.append(q)
                ps.append(p)
                scales.append(scale)

                rows.append({
                    "Criterion": criterion,
                    "Type": typ_display,
                    "Direction": direction_name,
                    "Preference": pref_name,
                    "q": q,
                    "p": p,
                    "Scale": (
                        scale
                        if typ == "Ordinal"
                        else "Continuous"
                    ),
                })

            config = pd.DataFrame(rows)

            st.divider()
            st.header("3. Execution")

            run = st.button(
                "▶ Run E-PROM",
                type="primary",
                use_container_width=True,
            )

        except Exception as exc:

            st.error(
                f"Error reading the workbook: {exc}"
            )


# ================================================================
# EXECUTION
# ================================================================

if run:

    if fahp_file is None:
        st.error(
            "Upload the FAHP-Express workbook."
        )
        st.stop()

    try:

        with st.spinner(
            "Running FAHP-Express + Fuzzy PROMETHEE..."
        ):

            data = preview

            fahp = read_fahp_workbook(
                normalize_fahp_workbook(
                    save_uploaded(fahp_file),
                    data["criteria"],
                ),
                n_criteria=len(
                    data["criteria"]
                ),
            )

            result = run_e_prom(
                data["evaluations"],
                fahp["w_fuzzy_individual"],
                types,
                dirs,
                prefs,
                qs,
                ps,
                scales,
                precision_levels=fahp.get("precision_levels"),
            )

            ranking = result_df(
                data,
                result,
            )

            export_config = config.rename(columns={
                "Criterio": "Criterion",
                "Critério": "Criterion",
                "Tipo": "Type",
                "Direção": "Direction",
                "Função": "Preference",
                "q": "q",
                "p": "p",
                "Escala": "Scale",
            }).copy()
            xlsx = excel_bytes(
                data,
                result,
                export_config,
                fahp,
            )

        st.session_state.update(
            data=data,
            fahp=fahp,
            result=result,
            config=config,
            ranking=ranking,
            xlsx=xlsx,
        )

        st.success(
            "E-PROM completed successfully."
        )

    except Exception as exc:
        st.exception(exc)
        st.stop()


# ================================================================
# RESULTS
# ================================================================

if "result" in st.session_state:

    data = st.session_state["data"]
    fahp = st.session_state["fahp"]
    result = st.session_state["result"]
    config = st.session_state["config"]
    ranking = st.session_state["ranking"]

    st.divider()

    a, b, c = st.columns(3)

    a.metric(
        "Decision makers",
        data["n_experts"],
    )

    b.metric(
        "Alternatives",
        data["n_requirements"],
    )

    c.metric(
        "Criteria",
        data["n_criteria"],
    )

    tab1, tab2, tab3, tab4, tab5 = st.tabs(
        [
            "🏆 Ranking",
            "📈 PROMETHEE",
            "👥 Decision Makers",
            "🎲 CPP / Monte Carlo",
            "📘 Guide",
        ]
    )

    with tab1:

        st.dataframe(
            ranking,
            use_container_width=True,
            hide_index=True,
        )

        st.download_button(
            "⬇️ Download results as Excel",
            data=st.session_state["xlsx"],
            file_name="E_PROM_Results.xlsx",
            mime=(
                "application/vnd.openxmlformats-officedocument."
                "spreadsheetml.sheet"
            ),
        )

    with tab2:

        st.subheader("Fuzzy criterion weights")
        st.caption(
            "Weights obtained by FAHP-Express. Fuzzy numbers are preserved during Fuzzy PROMETHEE and are not defuzzified at this stage."
        )
        fuzzy_weights_df = pd.DataFrame(
            np.asarray(fahp["w_fuzzy_group"], dtype=float),
            index=data["criteria"],
            columns=["a", "b", "c", "d"],
        )
        fuzzy_weights_df.index.name = "Criterion"
        st.dataframe(
            fuzzy_weights_df.style.format("{:.6f}"),
            use_container_width=True,
        )

        st.subheader("Attribute precision")
        precision_matrix = fahp.get(
            "precision_levels",
            [["Alta"] * len(data["criteria"]) for _ in fahp.get("expert_names", [])],
        )
        precision_df = pd.DataFrame(
            precision_matrix,
            index=fahp.get("expert_names", []),
            columns=data["criteria"],
        )
        precision_df.index.name = "Decision maker"
        st.dataframe(precision_df, use_container_width=True)
        st.caption(
            "Precision is defined individually for each decision maker and criterion. "
            "High = lower uncertainty, Medium = intermediate, and Low = higher "
            "uncertainty. For continuous criteria, the original unit is preserved; for "
            "ordinal criteria, fuzzy width is adjusted in scale units, "
            "respecting their limits."
        )

        st.plotly_chart(
            plot_weights(
                fahp,
                data["criteria"],
            ),
            use_container_width=True,
        )

        st.plotly_chart(
            plot_relation(
                data,
                result,
            ),
            use_container_width=True,
        )

        st.plotly_chart(
            plot_rank_heatmap(
                data,
                result,
            ),
            use_container_width=True,
        )

    with tab3:

        phi_df = pd.DataFrame(
            result["phi_net_experts"],
            index=data["ids"],
            columns=result["expert_names"],
        )

        phi_df.index.name = "Alternative"

        st.subheader(
            "Net Phi by decision maker"
        )

        st.dataframe(
            phi_df,
            use_container_width=True,
        )

        st.subheader(
            "Configuration"
        )

        st.dataframe(
            config,
            use_container_width=True,
            hide_index=True,
        )

    with tab4:

        st.subheader("CPP / Monte Carlo — E-PROM stability")

        st.markdown(
            "For **ordinal criteria**, uncertainty is obtained empirically "
            "from decision makers’ responses. For **continuous criteria**, "
            "enter the desired percentage variation."
        )
        st.info(
            "In E-PROM, decision makers’ flows remain fuzzy during "
            "aggregation. Defuzzification occurs only after fuzzy aggregation "
            "in each simulation."
        )

        continuous_idx = [
            i for i, typ in enumerate(config["Type"])
            if typ == "Continuous"
        ]

        variation = []
        for k, criterion in enumerate(data["criteria"]):
            if k in continuous_idx:
                default = 20.0
                variation.append(
                    st.number_input(
                        f"Variation for {criterion} (%)",
                        min_value=0.0,
                        max_value=1000.0,
                        value=default,
                        step=5.0,
                        key=f"mc_var_{k}",
                        help="Applies ± percentage variation to each observed decision-maker value.",
                    )
                )
            else:
                variation.append(0.0)

        col1, col2, col3 = st.columns(3)
        with col1:
            n_mc = st.number_input(
                "Number of simulations",
                min_value=100,
                max_value=500000,
                value=2000,
                step=100,
                key="mc_n",
            )
        with col2:
            seed = st.number_input(
                "Seed",
                min_value=0,
                value=42,
                step=1,
                key="mc_seed",
            )
        with col3:
            st.metric(
                "Continuous criteria",
                len(continuous_idx),
            )

        if st.button(
            "▶ Run CPP / Monte Carlo",
            type="primary",
            key="run_mc",
        ):

            # Usa as linhas de referência individuais fornecidas pelos
            # especialistas no FAHP-Express. Mantém compatibilidade com
            # versões anteriores do data_loader que não armazenavam
            # explicitamente `reference_rows`.
            if "reference_rows" in fahp:
                reference_rows = np.asarray(
                    fahp["reference_rows"],
                    dtype=float,
                )
            else:
                # Fallback: A_individual preserva as razões v_j/v_i
                # usadas para reconstruir o FAHP-Express. Como uma
                # constante comum não altera as razões, a primeira linha
                # da matriz pode ser usada como uma representação
                # equivalente dos valores de referência para o MC.
                A_individual = fahp.get("A_individual")
                if A_individual is None:
                    st.error(
                        "Could not recover the FAHP-Express reference inputs "
                        "from FAHP-Express for CPP/Monte Carlo. Re-run "
                        "E-PROM with the updated FAHP-Express workbook."
                    )
                    st.stop()

                reference_rows = np.asarray(
                    [np.asarray(A, dtype=float)[0, :] for A in A_individual],
                    dtype=float,
                )

            mc = run_e_prom_monte_carlo(
                data["evaluations"],
                reference_rows,
                types,
                dirs,
                prefs,
                qs,
                ps,
                scales,
                variation,
                n_mc=int(n_mc),
                seed=int(seed),
                precision_levels=fahp.get("precision_levels"),
            )

            st.session_state["mc"] = mc

        if "mc" in st.session_state:

            mc = st.session_state["mc"]

            st.success(
                f"Monte Carlo completed: {mc['n_mc']:,} simulations."
            )

            st.dataframe(
                mc_summary_df(data, mc),
                use_container_width=True,
                hide_index=True,
            )

            st.plotly_chart(
                plot_mc_rank_acceptability(data, mc),
                use_container_width=True,
            )

            st.plotly_chart(
                plot_mc_outworking(data, mc),
                use_container_width=True,
            )

            w_df = pd.DataFrame(
                mc["W_samples"],
                columns=data["criteria"],
            )

            st.subheader("Weight distribution")
            st.dataframe(
                w_df.describe().T[
                    ["mean", "std", "min", "max"]
                ],
                use_container_width=True,
            )

            st.download_button(
                "⬇️ Download Monte Carlo results",
                data=pd.DataFrame({
                    "Alternative": data["ids"],
                    "Mean position": mc["mean_rank"],
                    "Position SD": mc["std_rank"],
                    "1st-place frequency (%)": 100 * mc["freq_top1"],
                    "Mean Phi": mc["mean_phi"],
                    "Phi SD": mc["std_phi"],
                }).to_csv(index=False).encode("utf-8-sig"),
                file_name="E_PROM_CPP_MonteCarlo.csv",
                mime="text/csv",
            )

    with tab5:
        st.markdown(GUIDE_TEXT)