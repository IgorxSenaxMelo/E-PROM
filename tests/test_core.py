import sys
from pathlib import Path

import numpy as np
import pytest
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from core.data_loader import read_expert_workbook, read_fahp_workbook
from core.e_prom_core import (
    TFN_EVAL_5, TFN_EVAL_7, TFN_EVAL_9, build_fuzzy_matrix, continuous_to_fuzzy,
    fuzzify_ordinal, run_e_prom,
)
from core.fuzzy_numbers import defuzz_coa
from core.monte_carlo import run_e_prom_monte_carlo
from core.promethee import promethee_i_relation, rankdata_average

T = ROOT / "templates"
EVAL = T / "E_PROM_Evaluations_Template.xlsx"
FAHP = T / "E_PROM_FAHP_Express_Template.xlsx"

TYPES = ["Ordinal"] * 7 + ["Continuous"]
DIRS = [1] * 7 + [-1]
PREFS = [1] * 7 + [3]
QS = [0.0] * 8
PS = [0.0] * 7 + [10000.0]
SCALES = [9] * 7 + [0]


def load(eval_path=EVAL, fahp_path=FAHP):
    data = read_expert_workbook(eval_path)
    fahp = read_fahp_workbook(fahp_path, data["criteria"], data["expert_names"])
    return data, fahp


def test_templates_load_and_run():
    data, fahp = load()
    res = run_e_prom(data["evaluations"], fahp["w_fuzzy_individual"], TYPES, DIRS,
                     PREFS, QS, PS, SCALES, precision_levels=fahp["precision_levels"])
    assert list(res["rank_e_prom"]) == [2, 1, 3, 4]


def test_ordinal_scales_are_consistent():
    # Medium-precision trapezoid of every interior level is centred on its value.
    for table in (TFN_EVAL_7, TFN_EVAL_9):
        for v in range(2, len(table)):
            assert defuzz_coa(table[v - 1]) == pytest.approx(v, abs=1e-9)


def test_vectorized_fuzzification_matches_scalar():
    rng = np.random.default_rng(0)
    ordv = rng.integers(1, 10, 20).astype(float)
    cont = rng.uniform(1, 1e5, 20)
    for prec in ("High", "Medium", "Low", "Alta", "Média", "Baixa"):
        data = np.column_stack([ordv, cont])
        F = build_fuzzy_matrix(data, ["Ordinal", "Continuous"], [9, 0], [prec, prec])
        for i in range(20):
            assert np.allclose(F[i, 0], fuzzify_ordinal(ordv[i], 9, prec))
            assert np.allclose(F[i, 1], continuous_to_fuzzy(cont[i], prec))


def test_blank_precision_defaults_to_high(tmp_path):
    wb = load_workbook(FAHP)
    for ws in wb.worksheets:
        if ws.title in ("IGOR", "LUMA"):
            for c in range(2, 10):
                ws.cell(5, c).value = None
    p = tmp_path / "blank.xlsx"
    wb.save(p)
    data, fahp = load(fahp_path=p)
    assert all(x == "High" for row in fahp["precision_levels"] for x in row)
    run_e_prom(data["evaluations"], fahp["w_fuzzy_individual"], TYPES, DIRS, PREFS,
               QS, PS, SCALES, precision_levels=fahp["precision_levels"])


def test_promethee_i_weak_strict_rule():
    assert promethee_i_relation(0.5, 0.2, 0.3, 0.2) == "P"
    assert promethee_i_relation(0.3, 0.2, 0.5, 0.2) == "P-"
    assert promethee_i_relation(0.5, 0.1, 0.3, 0.2) == "P"
    assert promethee_i_relation(0.5, 0.3, 0.3, 0.2) == "R"
    assert promethee_i_relation(0.4, 0.2, 0.4, 0.2) == "I"


def test_ties_use_tolerance():
    assert list(rankdata_average([1.0, 1.0 + 1e-15, 0.5])) == [2.5, 2.5, 1.0]


def test_monte_carlo_precision_is_not_transposed_for_square_case():
    """n experts == n criteria used to silently transpose the precision."""
    rng = np.random.default_rng(0)
    n = 3
    ev = {f"E{e}": rng.integers(1, 10, (4, n)).astype(float) for e in range(n)}
    ref = np.array([[1, 2, 3], [1, 3, 2], [1, 2, 2]], float)
    prec = [["High", "Medium", "Low"], ["High", "High", "High"], ["Low", "Low", "Low"]]
    kw = dict(types=["Ordinal"] * n, directions=[1] * n, preference_types=[1] * n,
              q_vals=[0] * n, p_vals=[0] * n, scales=[9] * n,
              continuous_variation_pct=[0] * n, n_mc=50, seed=3)
    a = run_e_prom_monte_carlo(ev, ref, precision_levels=prec, **kw)
    b = run_e_prom_monte_carlo(ev, ref, precision_levels=np.array(prec, dtype=object).T.tolist(), **kw)
    assert not np.allclose(a["phi_mc"], b["phi_mc"])
    with pytest.raises(ValueError):
        run_e_prom_monte_carlo(
            {k: v[:, :2] for k, v in ev.items()}, ref[:, :2], types=["Ordinal"] * 2,
            directions=[1] * 2, preference_types=[1] * 2, q_vals=[0] * 2, p_vals=[0] * 2,
            scales=[9] * 2, continuous_variation_pct=[0] * 2, n_mc=5,
            precision_levels=np.array(prec, dtype=object)[:2, :2].T.tolist())


def test_monte_carlo_outputs_are_consistent():
    data, fahp = load()
    for mode in ("dm_bootstrap", "independent"):
        mc = run_e_prom_monte_carlo(
            data["evaluations"], fahp["reference_rows"], TYPES, DIRS, PREFS, QS, PS,
            SCALES, [0] * 7 + [20.0], n_mc=200, seed=1,
            precision_levels=fahp["precision_levels"], weight_sampling=mode)
        assert np.allclose(mc["rank_acceptability"].sum(axis=1), 100)
        assert mc["freq_top1"].sum() == pytest.approx(1.0)
        assert np.allclose(mc["W_samples"].sum(axis=1), 1)


def test_portuguese_workbooks_are_accepted(tmp_path):
    ev = load_workbook(EVAL)
    for ws in ev.worksheets:
        if ws.title in ("IGOR", "LUMA"):
            ws["B1"].value = "Descrição"
    ev.save(tmp_path / "ev_pt.xlsx")

    pt_terms = {
        "equal importance": "Igual importância",
        "weakly more important": "Fracamente mais importante",
        "moderately more important": "Moderadamente mais importante",
        "moderately to strongly more important": "Moderadamente a fortemente mais importante",
        "strongly more important": "Fortemente mais importante",
        "strongly to very strongly more important": "Fortemente a muito fortemente mais importante",
        "very strongly more important": "Muito fortemente mais importante",
        "very strongly to extremely more important": "Muito fortemente a extremamente mais importante",
        "extremely more important": "Extremamente mais importante",
    }
    pt_prec = {"High": "Alta", "Medium": "Média", "Low": "Baixa"}
    fh = load_workbook(FAHP)
    for ws in fh.worksheets:
        if ws.title in ("IGOR", "LUMA"):
            ws["A3"].value = "Critérios"
            ws["A4"].value = "Comparação com referência"
            ws["A5"].value = "Precisão do atributo"
            for c in range(2, 10):
                v = ws.cell(4, c).value
                ws.cell(4, c).value = pt_terms[v.lower()]
                ws.cell(5, c).value = pt_prec[ws.cell(5, c).value]
    fh.save(tmp_path / "fahp_pt.xlsx")

    d_en, f_en = load()
    d_pt, f_pt = load(tmp_path / "ev_pt.xlsx", tmp_path / "fahp_pt.xlsx")
    assert np.allclose(f_en["w_fuzzy_individual"], f_pt["w_fuzzy_individual"])
    assert f_en["precision_levels"] == f_pt["precision_levels"]


def test_order_of_experts_and_columns_does_not_matter(tmp_path):
    """FAHP sheets in a different order are paired by name."""
    fh = load_workbook(FAHP)
    fh.move_sheet("LUMA", offset=-1)
    fh.save(tmp_path / "reordered.xlsx")
    d, f = load(fahp_path=tmp_path / "reordered.xlsx")
    d0, f0 = load()
    assert f["expert_names"] == d["expert_names"]
    assert np.allclose(f["w_fuzzy_individual"], f0["w_fuzzy_individual"])
    assert f["warnings"] == []


def test_missing_guide_sheet_does_not_drop_a_decision_maker(tmp_path):
    ev = load_workbook(EVAL)
    del ev["Guide"]
    ev.save(tmp_path / "noguide.xlsx")
    d = read_expert_workbook(tmp_path / "noguide.xlsx")
    assert d["n_experts"] == 2
