"""Workbook readers for E-PROM.

Evaluation workbook
    - one sheet per decision maker (DM); the sheet name is the DM name;
    - header row with ``ID`` and ``Description`` in columns A and B, followed
      by one column per criterion;
    - every DM must contain the same alternative IDs and criterion names (the
      order of rows and columns may differ; they are aligned by name);
    - sheets named Guide / _Lists / Instructions ... are ignored.

FAHP-Express workbook
    - one sheet per DM (the DM names must match the evaluation workbook);
    - a ``Criteria`` header row: the FIRST criterion is that DM's reference;
    - a ``Comparison with reference`` row: how many times the reference is
      more important than each criterion (linguistic term or number 1-9);
    - an ``Attribute precision`` row: High / Medium / Low per criterion.

English and Portuguese labels are both accepted.
"""

from __future__ import annotations
from pathlib import Path

import numpy as np
from openpyxl import load_workbook

from .fahp import group_weights_from_reference_rows
from .labels import (
    is_support_sheet,
    match_key,
    normalize_precision,
    parse_importance,
)

_ID_KEYS = {"id"}
_DESCRIPTION_KEYS = {"description", "descricao"}
_CRITERIA_KEYS = {"criteria", "criterios"}
_COMPARISON_KEYS = {"comparison with reference", "comparacao com referencia"}
_PRECISION_KEYS = {"attribute precision", "precisao do atributo"}


def _text(value):
    return "" if value is None else str(value).strip()


def _is_empty(value):
    return value is None or (isinstance(value, str) and not value.strip())


def _to_number(x):
    if _is_empty(x):
        raise ValueError("Empty value found in the workbook.")

    if isinstance(x, (int, float, np.integer, np.floating)):
        value = float(x)
    else:
        try:
            value = float(str(x).strip().replace(",", "."))
        except ValueError:
            raise ValueError(
                f"Non-numeric value found in the evaluation matrix: {x!r}"
            ) from None

    if not np.isfinite(value):
        raise ValueError("NaN/inf value found in the workbook.")
    return value


def _decision_maker_sheets(wb):
    return [ws for ws in wb.worksheets if not is_support_sheet(ws.title)]


# --------------------------------------------------------------------------
# Evaluation workbook
# --------------------------------------------------------------------------

def _find_evaluation_header(ws, max_scan_rows=20):
    for r in range(1, min(ws.max_row, max_scan_rows) + 1):
        if (
            match_key(ws.cell(r, 1).value) in _ID_KEYS
            and match_key(ws.cell(r, 2).value) in _DESCRIPTION_KEYS
        ):
            return r
    return None


def _read_evaluation_sheet(ws):
    header_row = _find_evaluation_header(ws)
    if header_row is None:
        raise ValueError(
            f"Sheet '{ws.title}': could not find the header row. Expected "
            "'ID' in column A and 'Description' in column B."
        )

    criteria = []
    columns = []
    for c in range(3, ws.max_column + 1):
        name = _text(ws.cell(header_row, c).value)
        if name:
            criteria.append(name)
            columns.append(c)

    if not criteria:
        raise ValueError(
            f"Sheet '{ws.title}' needs an ID, a description and at least one criterion."
        )

    keys = [match_key(c) for c in criteria]
    if len(set(keys)) != len(keys):
        raise ValueError(f"Sheet '{ws.title}': duplicated criterion names.")

    ids, descriptions, values = [], [], []
    for r in range(header_row + 1, ws.max_row + 1):
        row_cells = [ws.cell(r, c).value for c in range(1, ws.max_column + 1)]
        if all(_is_empty(v) for v in row_cells):
            break                                  # first blank row ends the data
        alt_id = _text(ws.cell(r, 1).value)
        if not alt_id:
            raise ValueError(f"Sheet '{ws.title}', row {r}: missing alternative ID.")
        try:
            row_values = [_to_number(ws.cell(r, c).value) for c in columns]
        except ValueError as exc:
            raise ValueError(f"Sheet '{ws.title}', alternative {alt_id}: {exc}") from None
        ids.append(alt_id)
        descriptions.append(_text(ws.cell(r, 2).value))
        values.append(row_values)

    if len(ids) < 2:
        raise ValueError(f"Sheet '{ws.title}': at least two alternatives are required.")
    if len(set(ids)) != len(ids):
        raise ValueError(f"Sheet '{ws.title}': duplicated alternative IDs.")

    return {
        "ids": ids,
        "descriptions": descriptions,
        "criteria": criteria,
        "keys": keys,
        "values": np.asarray(values, dtype=float),
    }


def read_expert_workbook(path):
    """Read the evaluation workbook (see module docstring)."""
    wb = load_workbook(Path(path), data_only=True)
    sheets = _decision_maker_sheets(wb)

    if not sheets:
        raise ValueError(
            "The workbook has no decision-maker sheets (sheets named Guide, "
            "_Lists or Instructions are ignored)."
        )

    evaluations = {}
    ref = None

    for ws in sheets:
        parsed = _read_evaluation_sheet(ws)

        if ref is None:
            ref = parsed
            evaluations[ws.title] = parsed["values"]
            continue

        if set(parsed["keys"]) != set(ref["keys"]):
            missing = [c for c, k in zip(ref["criteria"], ref["keys"])
                       if k not in parsed["keys"]]
            extra = [c for c, k in zip(parsed["criteria"], parsed["keys"])
                     if k not in ref["keys"]]
            details = []
            if missing:
                details.append("missing: " + ", ".join(missing))
            if extra:
                details.append("extra: " + ", ".join(extra))
            raise ValueError(
                f"Sheet '{ws.title}' has different criteria from the first "
                f"decision maker ({'; '.join(details)})."
            )

        id_pos = {a: i for i, a in enumerate(parsed["ids"])}
        if set(id_pos) != set(ref["ids"]):
            raise ValueError(
                f"Sheet '{ws.title}' has different alternative IDs from the "
                "first decision maker."
            )

        col_pos = {k: i for i, k in enumerate(parsed["keys"])}
        row_order = [id_pos[a] for a in ref["ids"]]
        col_order = [col_pos[k] for k in ref["keys"]]
        evaluations[ws.title] = parsed["values"][np.ix_(row_order, col_order)]

    names = list(evaluations.keys())
    return {
        "sheets": names,
        "expert_names": names,
        "ids": ref["ids"],
        "descriptions": ref["descriptions"],
        "criteria": ref["criteria"],
        "evaluations": evaluations,
        "n_experts": len(names),
        "n_requirements": len(ref["ids"]),
        "n_criteria": len(ref["criteria"]),
    }


# --------------------------------------------------------------------------
# FAHP-Express workbook
# --------------------------------------------------------------------------

def _find_fahp_header(ws, max_scan_rows=30):
    for r in range(1, min(ws.max_row, max_scan_rows) + 1):
        if match_key(ws.cell(r, 1).value) not in _CRITERIA_KEYS:
            continue
        filled = sum(
            1 for c in range(2, ws.max_column + 1)
            if not _is_empty(ws.cell(r, c).value)
        )
        if filled >= 2:
            return r
    return None


def _find_labelled_row(ws, header_row, keys, fallback_offset):
    for r in range(header_row + 1, min(ws.max_row, header_row + 10) + 1):
        if match_key(ws.cell(r, 1).value) in keys:
            return r
    fallback = header_row + fallback_offset
    return fallback if fallback <= ws.max_row else None


def _read_fahp_sheet(ws, criteria):
    header_row = _find_fahp_header(ws)
    if header_row is None:
        raise ValueError(
            f"FAHP sheet '{ws.title}': could not find the 'Criteria' header row."
        )

    columns = {}
    order = []
    for c in range(2, ws.max_column + 1):
        name = _text(ws.cell(header_row, c).value)
        if not name:
            continue
        key = match_key(name)
        if key in columns:
            raise ValueError(f"FAHP sheet '{ws.title}': duplicated criterion '{name}'.")
        columns[key] = c
        order.append(name)

    wanted = {match_key(c): c for c in criteria}
    missing = [name for key, name in wanted.items() if key not in columns]
    extra = [name for name in order if match_key(name) not in wanted]
    if missing or extra:
        details = []
        if missing:
            details.append("missing: " + ", ".join(missing))
        if extra:
            details.append("extra: " + ", ".join(extra))
        raise ValueError(
            f"FAHP sheet '{ws.title}' has different criteria from the "
            f"evaluation workbook ({'; '.join(details)})."
        )

    comparison_row = _find_labelled_row(ws, header_row, _COMPARISON_KEYS, 1)
    if comparison_row is None:
        raise ValueError(
            f"FAHP sheet '{ws.title}': could not find 'Comparison with reference'."
        )
    precision_row = _find_labelled_row(ws, header_row, _PRECISION_KEYS, 2)

    reference_key = match_key(order[0])
    reference_name = order[0]

    row = []
    precision = []
    for criterion in criteria:
        col = columns[match_key(criterion)]
        try:
            row.append(parse_importance(ws.cell(comparison_row, col).value))
        except ValueError as exc:
            raise ValueError(
                f"FAHP sheet '{ws.title}', criterion '{criterion}': {exc}"
            ) from None
        try:
            precision.append(
                normalize_precision(
                    ws.cell(precision_row, col).value if precision_row else None
                )
            )
        except ValueError as exc:
            raise ValueError(
                f"FAHP sheet '{ws.title}', criterion '{criterion}': {exc}"
            ) from None

    ref_idx = [match_key(c) for c in criteria].index(reference_key)
    if not np.isclose(row[ref_idx], 1.0, atol=1e-9):
        raise ValueError(
            f"FAHP sheet '{ws.title}': the first criterion ('{reference_name}') "
            "is the reference and must be 'Equal importance' (1)."
        )
    if any(v < 1.0 - 1e-9 for v in row):
        raise ValueError(
            f"FAHP sheet '{ws.title}': the reference criterion must be the most "
            "important one, so every other value must be >= 1."
        )

    return row, precision, reference_name


def read_fahp_workbook(path, criteria, expert_names=None):
    """Read the FAHP-Express workbook.

    ``criteria`` are the canonical criterion names (from the evaluation
    workbook). Values are returned in that canonical order, whatever order and
    reference criterion each DM used. If ``expert_names`` is given, the DMs are
    matched to the evaluation workbook by sheet name (falling back to sheet
    position, with a warning, when the names differ).
    """
    wb = load_workbook(Path(path), data_only=True)
    sheets = _decision_maker_sheets(wb)

    if not sheets:
        raise ValueError("The FAHP-Express workbook has no decision-maker sheets.")
    if len(criteria) < 2:
        raise ValueError("FAHP-Express requires at least 2 criteria.")

    warnings = []

    if expert_names is not None:
        expert_names = list(expert_names)
        if len(sheets) != len(expert_names):
            raise ValueError(
                f"The FAHP-Express workbook has {len(sheets)} decision-maker "
                f"sheets but the evaluation workbook has {len(expert_names)}."
            )
        by_key = {match_key(ws.title): ws for ws in sheets}
        if set(by_key) == {match_key(n) for n in expert_names} and len(by_key) == len(sheets):
            sheets = [by_key[match_key(n)] for n in expert_names]
        else:
            warnings.append(
                "The decision-maker sheet names differ between the two workbooks "
                f"({[ws.title for ws in sheets]} vs {expert_names}); they were "
                "paired by sheet position. Rename the sheets identically to pair "
                "them by name."
            )

    rows, precision, references, names = [], [], [], []
    for ws in sheets:
        row, prec, ref_name = _read_fahp_sheet(ws, criteria)
        rows.append(row)
        precision.append(prec)
        references.append(ref_name)
        names.append(ws.title)

    result = group_weights_from_reference_rows(rows)
    result["reference_rows"] = np.asarray(rows, dtype=float)      # [DM, criteria]
    result["precision_levels"] = precision                         # [DM][criteria]
    result["reference_criteria"] = references
    result["sheets"] = names
    result["expert_names"] = names
    result["warnings"] = warnings
    return result
